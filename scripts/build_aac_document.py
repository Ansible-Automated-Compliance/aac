#!/usr/bin/env python3
"""Build an AAC house-style DOCX and PDF from a Markdown source.

This exists because the build scripts for this kept living in /tmp and kept
disappearing with it -- the method was recorded in memory, the code was not.
Every AAC generated document needs the same two artefacts in the same style,
so the method belongs in the repo.

## How it works, and why it is a hybrid

  DOCX   pandoc --reference-doc, then a python-docx post-pass.

         The house template's *style definitions* already carry the AAC
         palette (Heading 1 = CC0000/14pt, Heading 2/3 = 006EA6, Heading 4 =
         555555), so pandoc's output renders correctly in Word even though
         the runs carry no direct formatting. Inspecting a pandoc DOCX with
         python-docx shows `color=None` on every run, which looks broken and
         is not -- the colour is inherited from the named style.

         What the template does NOT carry is table shading: its `Table`
         style defines no header fill, so pandoc tables come out flat. The
         post-pass supplies the navy header, white bold header text,
         alternating row fill, and a repeating header row across pages.

  PDF    pandoc -> standalone HTML with the house stylesheet, then Chrome
         headless. No LaTeX, no LibreOffice, and weasyprint is missing
         libgobject on this machine, so Chrome is the only engine available.

## Usage

    build_aac_document.py docs/m365/FOO.md
    build_aac_document.py docs/m365/FOO.md --no-pdf
    build_aac_document.py docs/m365/FOO.md --outdir ~/Downloads

Version and date are read from the document's own header lines:

    **Version:** v1.1
    **Date:** 2026-08-25

so the filenames cannot drift from the content. Four artefacts are written:

    <Name>-v<X.Y>.docx        <Name>-v<X.Y>-<YYYY-MM-DD>.docx
    <Name>-v<X.Y>.md          <Name>-v<X.Y>-<YYYY-MM-DD>.md
    (+ <Name>-v<X.Y>.pdf      <Name>-v<X.Y>-<YYYY-MM-DD>.pdf  unless --no-pdf)

Existing files are never overwritten without --force: regenerating over a
previous version has destroyed originals before.
"""
from __future__ import annotations

import argparse
import pathlib
import re
import shutil
import subprocess
import sys

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

# House palette. Mirrors ~/Downloads/aac_house.css -- keep the two in step.
HEADER_FILL = "1A1A2E"   # dark navy
ROW_ALT_FILL = "F2F2F2"  # alternating body rows
HEADER_TEXT = RGBColor(0xFF, 0xFF, 0xFF)

TEMPLATE = pathlib.Path.home() / "Downloads" / "AAC_Document_Template.docx"
STYLESHEET = pathlib.Path.home() / "Downloads" / "aac_house.css"
CHROME = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"

VERSION_RE = re.compile(r"^\*\*Version:\*\*\s*v?([0-9]+\.[0-9]+)", re.M)
DATE_RE = re.compile(r"^\*\*Date:\*\*\s*(\d{4}-\d{2}-\d{2})", re.M)


def read_metadata(src: pathlib.Path) -> tuple[str, str]:
    """Take version and date from the document rather than the command line.

    A filename that disagrees with the version inside the file is worse than
    no filename convention at all -- it is silently wrong.
    """
    text = src.read_text(encoding="utf-8")
    version = VERSION_RE.search(text)
    date = DATE_RE.search(text)
    if not version:
        sys.exit(f"{src}: no '**Version:** vX.Y' line found")
    if not date:
        sys.exit(f"{src}: no '**Date:** YYYY-MM-DD' line found")
    return version.group(1), date.group(1)


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    for existing in tc_pr.findall(qn("w:shd")):
        tc_pr.remove(existing)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def repeat_header_row(row) -> None:
    """Mark the row as a header so Word repeats it across page breaks.

    Without this a table split over two pages loses its column labels on the
    second, which is exactly where a reader most needs them.
    """
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def style_tables(doc: Document) -> int:
    """Supply the table styling the reference template does not define."""
    for table in doc.tables:
        if not table.rows:
            continue
        for cell in table.rows[0].cells:
            shade(cell, HEADER_FILL)
            for para in cell.paragraphs:
                for run in para.runs:
                    run.font.bold = True
                    run.font.color.rgb = HEADER_TEXT
        repeat_header_row(table.rows[0])
        # Banding starts on the first body row, so the first row under the
        # navy header stays white and the stripe reads as intentional.
        for index, row in enumerate(table.rows[1:]):
            if index % 2 == 1:
                for cell in row.cells:
                    shade(cell, ROW_ALT_FILL)
    return len(doc.tables)


def build_docx(src: pathlib.Path, dest: pathlib.Path, template: pathlib.Path) -> int:
    cmd = ["pandoc", str(src)]
    if template.exists():
        cmd += ["--reference-doc", str(template)]
    else:
        print(f"  ! template missing ({template}); DOCX will use pandoc defaults")
    cmd += ["-o", str(dest)]
    subprocess.run(cmd, check=True)
    doc = Document(str(dest))
    count = style_tables(doc)
    doc.save(str(dest))
    return count


def build_pdf(src: pathlib.Path, dest: pathlib.Path, css: pathlib.Path,
              workdir: pathlib.Path) -> bool:
    if not shutil.which(CHROME) and not pathlib.Path(CHROME).exists():
        print(f"  ! Chrome not found at {CHROME}; skipping PDF")
        return False
    html = workdir / (dest.stem + ".html")
    cmd = ["pandoc", str(src), "-s", "-o", str(html)]
    if css.exists():
        cmd += ["--css", str(css), "--embed-resources"]
    else:
        print(f"  ! stylesheet missing ({css}); PDF will be unstyled")
    subprocess.run(cmd, check=True)
    subprocess.run(
        [CHROME, "--headless", "--disable-gpu", f"--print-to-pdf={dest}",
         "--print-to-pdf-no-header", "--no-pdf-header-footer", html.as_uri()],
        check=True, capture_output=True,
    )
    return dest.exists()


def place(path: pathlib.Path, force: bool) -> bool:
    """Refuse to clobber an existing artefact unless explicitly told to.

    Regenerating over a previous version has destroyed originals before;
    storage is cheap and a lost signed document is not.
    """
    if path.exists() and not force:
        print(f"  = {path.name} exists, left alone (use --force to replace)")
        return False
    return True


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("source", type=pathlib.Path, help="Markdown source")
    ap.add_argument("--outdir", type=pathlib.Path,
                    default=pathlib.Path.home() / "Downloads")
    ap.add_argument("--name", help="basename for artefacts (default: source stem)")
    ap.add_argument("--template", type=pathlib.Path, default=TEMPLATE)
    ap.add_argument("--css", type=pathlib.Path, default=STYLESHEET)
    ap.add_argument("--no-pdf", action="store_true",
                    help="skip the PDF (internal docs do not need one)")
    ap.add_argument("--force", action="store_true",
                    help="overwrite artefacts that already exist")
    args = ap.parse_args()

    src = args.source
    if not src.exists():
        sys.exit(f"no such file: {src}")
    if not shutil.which("pandoc"):
        sys.exit("pandoc not on PATH")

    version, date = read_metadata(src)
    name = args.name or src.stem
    out = args.outdir
    out.mkdir(parents=True, exist_ok=True)
    workdir = out / ".aac-build"
    workdir.mkdir(exist_ok=True)

    print(f"{name}  v{version}  {date}")
    written = []

    md_current = out / f"{name}-v{version}.md"
    md_dated = out / f"{name}-v{version}-{date}.md"
    for target in (md_current, md_dated):
        if place(target, args.force):
            target.write_text(src.read_text(encoding="utf-8"), encoding="utf-8")
            written.append(target)

    docx_current = out / f"{name}-v{version}.docx"
    docx_dated = out / f"{name}-v{version}-{date}.docx"
    if place(docx_current, args.force):
        tables = build_docx(src, docx_current, args.template)
        print(f"  + {docx_current.name}  ({tables} tables styled)")
        written.append(docx_current)
    if place(docx_dated, args.force) and docx_current.exists():
        shutil.copyfile(docx_current, docx_dated)
        written.append(docx_dated)

    if not args.no_pdf:
        pdf_current = out / f"{name}-v{version}.pdf"
        pdf_dated = out / f"{name}-v{version}-{date}.pdf"
        if place(pdf_current, args.force):
            if build_pdf(src, pdf_current, args.css, workdir):
                print(f"  + {pdf_current.name}")
                written.append(pdf_current)
        if place(pdf_dated, args.force) and pdf_current.exists():
            shutil.copyfile(pdf_current, pdf_dated)
            written.append(pdf_dated)

    shutil.rmtree(workdir, ignore_errors=True)
    print(f"\n{len(written)} artefact(s) written to {out}")
    for path in written:
        print(f"  {path.name}  ({path.stat().st_size:,} bytes)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
